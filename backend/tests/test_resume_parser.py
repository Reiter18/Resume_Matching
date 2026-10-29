import io
import zipfile

import pymupdf
import pytest
from docx import Document

from app import resume_parser
from app.resume_parser import ResumeError, clean_text, detect_language, extract_text, finalize


def make_pdf(text: str) -> bytes:
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), text)
    return doc.tobytes()


def make_docx(build) -> bytes:
    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_pdf_text_extracted():
    pdf = make_pdf("John Smith\nBackend developer\nSkills: Python, FastAPI, PostgreSQL, Docker")
    text = extract_text("cv.pdf", pdf)
    assert "FastAPI" in text and "Backend developer" in text


def test_two_column_pdf_reads_left_column_then_right():
    doc = pymupdf.open()
    page = doc.new_page()  
    left = [f"LEFT-{i} experience description line number {i}" for i in range(1, 5)]
    right = [f"RIGHT-{i} skills description line number {i}" for i in range(1, 5)]
    for i, (a, b) in enumerate(zip(left, right, strict=True)):
        page.insert_text((50, 100 + i * 60), a)
        page.insert_text((330, 100 + i * 60), b)
    text = extract_text("cv.pdf", doc.tobytes())
    order = [text.index(x.split()[0]) for x in left + right]
    assert order == sorted(order)  


def test_rows_with_right_aligned_dates_are_not_split_into_columns():
    doc = pymupdf.open()
    page = doc.new_page()
    rows = [("Backend developer at Acme", "2022 - 2024"), ("Junior developer at Beta", "2020 - 2022"),
            ("Intern at Gamma company", "2019 - 2020"), ("Student at University", "2015 - 2019")]  
    for i, (job, dates) in enumerate(rows):
        page.insert_text((50, 100 + i * 40), job)
        page.insert_text((470, 100 + i * 40), dates)
    lines = extract_text("cv.pdf", doc.tobytes()).split("\n")
    assert (
        lines[0].startswith("Backend developer") and "2022" in lines[1]
    )  


def test_single_column_order_is_top_to_bottom():
    doc = pymupdf.open()
    page = doc.new_page()
    for i, line in enumerate(["FIRST block about education", "SECOND block about skills", "THIRD block about work"]):
        page.insert_textbox(pymupdf.Rect(72, 80 + i * 100, 520, 140 + i * 100), line)
    text = extract_text("cv.pdf", doc.tobytes())
    assert text.index("FIRST") < text.index("SECOND") < text.index("THIRD")


def test_repeated_headers_and_footers_are_removed():
    words = ["Alpha", "Bravo", "Charlie", "Delta"]
    doc = pymupdf.open()
    for n in range(1, 5):
        page = doc.new_page()
        page.insert_text((72, 40), "John Smith | Resume")  
        page.insert_text((72, 120), f"{words[n - 1]} content of this page has enough different words to keep")
        page.insert_text((72, 800), f"John Smith page {n} of 4")
    text = extract_text("cv.pdf", doc.tobytes())
    assert "John Smith" not in text
    for w in words:
        assert f"{w} content of this page" in text


def test_scanned_pdf_without_text_gives_clear_error():
    doc = pymupdf.open()
    doc.new_page().draw_rect(pymupdf.Rect(50, 50, 200, 200))
    with pytest.raises(ResumeError, match="скан"):
        extract_text("scan.pdf", doc.tobytes())


def test_password_protected_pdf_rejected():
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "secret resume text that is long enough")
    data = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="123", owner_pw="456")
    with pytest.raises(ResumeError, match="парол"):
        extract_text("cv.pdf", data)


def test_too_many_pages_rejected(monkeypatch):
    monkeypatch.setattr(resume_parser, "MAX_PAGES", 2)
    doc = pymupdf.open()
    for _ in range(3):
        doc.new_page().insert_text((72, 72), "page with some text content here")
    with pytest.raises(ResumeError, match="страниц"):
        extract_text("cv.pdf", doc.tobytes())



def test_docx_paragraphs_and_tables_keep_document_order():
    def build(doc):
        doc.add_paragraph("John Smith, Backend developer")
        table = doc.add_table(rows=3, cols=2)
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Skills", "Python, FastAPI"
        table.rows[1].cells[0].text, table.rows[1].cells[1].text = "Languages", "English, Russian"
        merged = table.rows[2].cells[0].merge(table.rows[2].cells[1])
        merged.text = "Merged footer row"
        doc.add_paragraph("References available on request, thank you")

    text = extract_text("cv.docx", make_docx(build))
    assert "Skills | Python, FastAPI" in text and "Languages | English, Russian" in text
    assert text.count("Merged footer row") == 1  
    assert text.index("John Smith") < text.index("Skills") < text.index("References")


def test_docx_detected_by_content_not_extension():
    data = make_docx(lambda d: d.add_paragraph("Backend developer with Python and Docker experience"))
    assert "Python" in extract_text("resume", data)


def test_zip_that_is_not_docx_and_zip_bomb_guard(monkeypatch):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "hi")
    with pytest.raises(ResumeError, match="не документ DOCX"):
        extract_text("cv.docx", buf.getvalue())
    monkeypatch.setattr(resume_parser, "MAX_UNZIPPED_BYTES", 10)
    big = make_docx(lambda d: d.add_paragraph("some text that is definitely long enough to matter"))
    with pytest.raises(ResumeError, match="слишком большой"):
        extract_text("cv.docx", big)


def test_corrupted_docx_and_old_doc_format():
    with pytest.raises(ResumeError, match="повреждён"):
        extract_text("cv.docx", b"PK\x03\x04 garbage")
    with pytest.raises(ResumeError, match=r"\.doc"):
        extract_text("cv.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 old word file")



def test_txt_utf8_and_cp1251():
    src = "Иван Петров\nPython-разработчик, опыт 2 года, FastAPI, PostgreSQL"
    assert extract_text("cv.txt", src.encode("utf-8")) == src
    assert extract_text("cv.txt", src.encode("cp1251")) == src


def test_broken_pdf_unknown_format_empty_and_binary():
    with pytest.raises(ResumeError):
        extract_text("cv.pdf", b"%PDF-1.4 this is not a real pdf")
    with pytest.raises(ResumeError, match="PDF, DOCX и TXT"):
        extract_text("cv.png", b"\x89PNG\r\n")
    with pytest.raises(ResumeError, match="пустой"):
        extract_text("cv.txt", b"")
    with pytest.raises(ResumeError, match="не похож"):
        extract_text("cv.txt", b"abc\x00\x01\x02def" * 10)


def test_length_limits(monkeypatch):
    with pytest.raises(ResumeError, match="скан"):
        finalize("too short")
    monkeypatch.setattr(resume_parser, "MAX_CHARS", 50)
    with pytest.raises(ResumeError, match="длинный"):
        finalize("word " * 40)



def test_clean_text_rules():
    raw = "ﬁne\u00a0text\r\n\r\n\r\n\r\n• Python\r\n● Docker\r\nразра-\nботчик\r\nСтр. 2 из 3\r\nfull-stack\x00 dev"
    out = clean_text(raw)
    assert "fine text" in out
    assert "- Python" in out and "- Docker" in out
    assert "разработчик" in out
    assert "Стр" not in out
    assert "full-stack dev" in out and "\x00" not in out
    assert "\n\n\n" not in out


def test_detect_language():
    assert detect_language("Python-разработчик с опытом работы в команде над веб-сервисами и базами данных") == "ru"
    assert detect_language("Backend developer with five years of experience building web services") == "en"
    assert detect_language("Разработчик Python, experienced backend developer, работал с FastAPI и Docker") == "mixed"
    assert detect_language("12345 !!!") == "unknown"
